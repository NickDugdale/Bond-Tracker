"""
validate.py  (Stage 4)
Two checks on the extraction, together giving an honest accuracy figure.

1. GROUNDING CHECK (automatic, every bond)
   For every number and date Claude extracted, search the original filing
   for it. If a value can't be found anywhere in the source, the model
   invented or garbled it. This proves "nothing made up", but NOT "right
   number from the right place" (e.g. a make-whole "T+15" is in the text
   too, but it isn't the pricing spread). That's what check 2 is for.

2. SPOT CHECK (you, ~20 minutes)
   A random sample of 8 deals (fixed seed, so it's reproducible). For each
   key field you see Claude's value next to the source line it came from,
   and mark it Y (correct) or N (wrong). Because you see the answer first,
   this is "manual verification", not blind labelling; describe it that way.

HOW TO RUN
    python validate.py grounding     -> automatic check, prints results
    python validate.py sample        -> creates data/validation/spot_check.xlsx
    (fill in the Y/N column, save, close Excel)
    python validate.py score         -> spot-check accuracy + final summary
"""
import argparse
import calendar
import json
import random
import re
import sys
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config

SEED = 42                # fixed seed = same sample every time
N_DEALS = 8              # deals in the spot check
MAX_BONDS_PER_DEAL = 3   # stops one 7-bond deal eating all your time
SPOT_FILE = config.VALIDATION_DIR / "spot_check.xlsx"
GROUNDING_FILE = config.VALIDATION_DIR / "grounding.json"
GROUNDING_FAILS = config.VALIDATION_DIR / "grounding_failures.csv"
RESULTS_FILE = config.VALIDATION_DIR / "results.json"

# The fields that drive the analysis and the weekly report
TRANCHE_FIELDS = ["principal_amount", "coupon_pct", "maturity_date",
                  "first_call_date", "spread_bps", "yield_pct", "issue_price_pct"]
DEAL_FIELDS = ["rating_moodys", "rating_sp", "rating_fitch"]
DATE_FIELDS = {"maturity_date", "first_call_date"}

# Where each field is normally labelled in a term sheet. Lines with these
# words are searched FIRST, so the evidence shown is the relevant line
# (e.g. the par call line, not the make-whole line with the same date).
LABEL_HINTS = {
    "principal_amount": r"principal|amount|size",
    "coupon_pct": r"coupon|interest rate",
    "maturity_date": r"maturity",
    "first_call_date": r"par call|on or after",
    "spread_bps": r"spread",
    "yield_pct": r"yield",
    "issue_price_pct": r"price",
    "rating_moodys": r"rating", "rating_sp": r"rating", "rating_fitch": r"rating",
}


# ======================================================================
# LOADING
# ======================================================================
def load_bond_docs():
    """Every extracted document that is a bond, with its source text."""
    docs = []
    for path in sorted(config.EXTRACTED_DIR.glob("*.json")):
        record = json.loads(path.read_text(encoding="utf-8"))
        if not record["data"].get("is_bond_pricing"):
            continue
        raw = config.RAW_DIR / record["source_file"]
        text = raw.read_text(encoding="utf-8", errors="ignore") if raw.suffix == ".txt" else ""
        record["lines"] = [l for l in text.splitlines()[4:] if l.strip()]   # skip our header
        record["url"] = text.splitlines()[0].replace("SOURCE:", "").strip() if text else ""
        docs.append(record)
    return docs


# ======================================================================
# HOW TO SPOT A VALUE IN THE TEXT
# Each value is turned into the patterns it might be written as, e.g.
# 500000000 -> "500,000,000" or "500 million"; 2036-09-15 -> "September 15, 2036".
# ======================================================================
def number_pattern(v):
    """5.25 matches '5.25', '5.250', '5.2500' but not '15.25' or '5.251'."""
    s = f"{float(v):.6f}".rstrip("0").rstrip(".")
    tail = r"0*" if "." in s else r"(?:\.0+)?"
    return rf"(?<![\d.]){re.escape(s)}{tail}(?!\d)"


def amount_patterns(v):
    v = float(v)
    patterns = [re.escape(f"{int(v):,}"), rf"(?<![\d,]){int(v)}(?!\d)"]
    for unit, div in (("million", 1e6), ("billion", 1e9)):
        n = f"{v / div:.4f}".rstrip("0").rstrip(".")
        patterns.append(rf"(?<![\d.]){re.escape(n)}\s*{unit}")
    return patterns


def date_patterns(v):
    d = date.fromisoformat(v)
    month = calendar.month_name[d.month][:3]          # 'Sep' also matches 'September', 'Sept.'
    return [rf"{month}[a-z]*\.?\s+0?{d.day},?\s+{d.year}",
            rf"(?<!\d)0?{d.day}\s+{month}[a-z]*\.?,?\s+{d.year}",
            rf"0?{d.month}/0?{d.day}/{d.year}",
            re.escape(v)]


def rating_pattern(v):
    core = re.sub(r"^\(P\)\s*", "", str(v)).split()[0]
    return [rf"(?<![A-Za-z]){re.escape(core)}(?![A-Za-z0-9+-])"]


def patterns_for(field, value):
    if field == "spread_bps":
        # A bare "15" appears everywhere (dates, "T+5"...), so a spread only
        # counts if written like one: "+85", "T+85", "85 bps", "85 basis points".
        n = number_pattern(value)
        return [rf"(?:\+|T\s*\+)\s*{n}", rf"{n}\s*(?:bps|bp\b|basis)"]
    if field == "principal_amount":
        return amount_patterns(value)
    if field in DATE_FIELDS:
        try:
            return date_patterns(value)
        except ValueError:
            return []
    if field.startswith("rating"):
        return rating_pattern(value)
    return [number_pattern(value)]


def find_evidence(field, value, lines):
    """Return the source line containing the value (preferring lines with
    the field's usual label), or None if it isn't anywhere in the filing."""
    hint = re.compile(LABEL_HINTS.get(field, "$^"), re.I)
    preferred = [l for l in lines if hint.search(l)]
    flags = re.I if field in DATE_FIELDS or field == "spread_bps" else 0
    for pool in (preferred, lines):
        for pat in patterns_for(field, value):
            regex = re.compile(pat, flags)
            for line in pool:
                if regex.search(line):
                    return line.strip()
    return None


# ======================================================================
# CHECK 1: GROUNDING
# ======================================================================
def grounding():
    docs = load_bond_docs()
    checks = []
    for rec in docs:
        d = rec["data"]
        values = [(f, d.get(f), "-") for f in DEAL_FIELDS]
        for i, t in enumerate(d.get("tranches") or [], start=1):
            values += [(f, t.get(f), i) for f in TRANCHE_FIELDS]
        for field, value, bond in values:
            if value is None:
                continue        # a null can't be searched for; the spot check covers nulls
            found = find_evidence(field, value, rec["lines"])
            checks.append({"source_file": rec["source_file"], "bond": bond, "field": field,
                           "value": value, "grounded": found is not None})

    df = pd.DataFrame(checks)
    per_field = (df.groupby("field")["grounded"].agg(values="count", grounded="sum")
                 .assign(pct=lambda x: (100 * x["grounded"] / x["values"]).round(1))
                 .sort_values("pct"))
    result = {"date": date.today().isoformat(), "documents": len(docs),
              "values_checked": int(len(df)), "grounded_pct": round(100 * df["grounded"].mean(), 1),
              "per_field": per_field.reset_index().to_dict("records")}
    config.VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    GROUNDING_FILE.write_text(json.dumps(result, indent=2, default=int), encoding="utf-8")
    df[~df["grounded"]].to_csv(GROUNDING_FAILS, index=False, encoding="utf-8-sig")

    print(f"Checked {len(df)} extracted values across {len(docs)} bond documents.")
    print(f"GROUNDED (found in the source filing): {result['grounded_pct']}%\n")
    print(per_field.to_string())
    print(f"\n{(~df['grounded']).sum()} values not found; listed in {GROUNDING_FAILS.name}. "
          "Open a few: a real error, or just written in an unusual way?")


# ======================================================================
# CHECK 2: SPOT CHECK
# ======================================================================
def fmt(v):
    if v is None:
        return "NOT STATED"
    if isinstance(v, (int, float)) and not isinstance(v, bool) and float(v).is_integer():
        return f"{int(v):,}"
    return str(v)


def sample():
    if SPOT_FILE.exists():
        sys.exit(f"{SPOT_FILE.name} already exists. Delete it first if you want a new sample.")
    docs = load_bond_docs()
    rng = random.Random(SEED)
    picked = sorted(rng.sample(docs, min(N_DEALS, len(docs))), key=lambda r: r["source_file"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Check"
    headers = ["source_file", "source_url", "bond", "field", "claude_value",
               "evidence_from_source", "correct (Y/N)", "note"]
    ws.append(headers)
    for rec in picked:
        d = rec["data"]
        tranches = list(enumerate(d.get("tranches") or [], start=1))
        if len(tranches) > MAX_BONDS_PER_DEAL:
            tranches = sorted(rng.sample(tranches, MAX_BONDS_PER_DEAL))
        for f in DEAL_FIELDS:
            add_row(ws, rec, "deal", f, d.get(f))
        for i, t in tranches:
            label = f"#{i}: {t.get('currency')} {fmt(t.get('principal_amount'))} due {t.get('maturity_date')}"
            for f in TRANCHE_FIELDS:
                add_row(ws, rec, label, f, t.get(f))

    # Y/N dropdown, header styling, widths
    dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{ws.max_row}")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
    for col, width in zip("ABCDEFGH", (30, 30, 34, 16, 18, 80, 14, 30)):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[5].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "E2"

    info = wb.create_sheet("Instructions")
    for line in [
        "For each row, decide: is claude_value what the filing says for that field?",
        "Mark Y or N in the 'correct' column. Use 'note' to say what the right value is if N.",
        "",
        "evidence_from_source = the line in the filing where that value appears.",
        "  Check it's the RIGHT line, e.g. the spread must come from 'Spread to Benchmark',",
        "  NOT from 'Make-Whole Call: T+15'. Right number, wrong line = N.",
        "If evidence is blank or claude_value is NOT STATED, open source_url and Ctrl+F to check.",
        "  NOT STATED is correct (Y) only if the filing really doesn't give that value.",
        "first_call_date should be the PAR call date ('Par Call: on or after...'), not the make-whole.",
        "",
        f"Sample: {len(picked)} deals drawn with random seed {SEED}; deals with more than "
        f"{MAX_BONDS_PER_DEAL} bonds had {MAX_BONDS_PER_DEAL} bonds picked at random. Created {date.today()}.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 110

    config.VALIDATION_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(SPOT_FILE)
    print(f"Created {SPOT_FILE.name}: {ws.max_row - 1} values to check from {len(picked)} deals.")
    print("Fill in the Y/N column, save, close Excel, then run: python validate.py score")


def add_row(ws, rec, bond, field, value):
    evidence = find_evidence(field, value, rec["lines"]) if value is not None else None
    ws.append([rec["source_file"], rec["url"], bond, field, fmt(value), evidence or "", "", ""])


def score():
    if not SPOT_FILE.exists():
        sys.exit("No spot_check.xlsx yet. Run: python validate.py sample")
    df = pd.read_excel(SPOT_FILE, sheet_name="Check", dtype=str, keep_default_na=False)
    df["mark"] = df["correct (Y/N)"].str.strip().str.upper().str[:1]
    blank = df[~df["mark"].isin(["Y", "N"])]
    if len(blank):
        sys.exit(f"{len(blank)} rows still need a Y or N (first: row {blank.index[0] + 2}).")

    df["correct"] = df["mark"] == "Y"
    df["is_null"] = df["claude_value"] == "NOT STATED"
    stated, nulls = df[~df["is_null"]], df[df["is_null"]]
    per_field = (df.groupby("field")["correct"].agg(checked="count", correct="sum")
                 .assign(pct=lambda x: (100 * x["correct"] / x["checked"]).round(1))
                 .sort_values("pct"))

    def pct(frame):
        return round(100 * frame["correct"].mean(), 1) if len(frame) else None

    results = {
        "date_scored": date.today().isoformat(), "model": config.MODEL, "seed": SEED,
        "spot_check": {"deals": int(df["source_file"].nunique()), "values_checked": int(len(df)),
                       "accuracy_pct": pct(df), "stated_values": int(len(stated)),
                       "stated_accuracy_pct": pct(stated), "nulls": int(len(nulls)),
                       "null_accuracy_pct": pct(nulls),
                       "per_field": per_field.reset_index().to_dict("records")},
        "grounding": json.loads(GROUNDING_FILE.read_text()) if GROUNDING_FILE.exists() else None,
    }
    RESULTS_FILE.write_text(json.dumps(results, indent=2, default=int), encoding="utf-8")

    s = results["spot_check"]
    print(f"SPOT CHECK: {s['accuracy_pct']}% correct across {s['values_checked']} values "
          f"from {s['deals']} random deals")
    print(f"  stated values: {s['stated_accuracy_pct']}% ({s['stated_values']}) | "
          f"'not stated': {s['null_accuracy_pct']}% ({s['nulls']})\n")
    print(per_field.to_string())
    wrong = df[~df["correct"]]
    if len(wrong):
        print("\nMarked wrong:")
        print(wrong[["source_file", "bond", "field", "claude_value", "note"]].to_string(index=False))
    g = results["grounding"]
    print("\n--- CV summary ---")
    if g:
        print(f"{g['grounded_pct']}% of {g['values_checked']} extracted values traced to the source "
              f"filing; key fields manually verified at {s['accuracy_pct']}% on a random sample "
              f"of {s['values_checked']} values.")
    else:
        print("Run 'python validate.py grounding' too, to complete the summary.")


def main():
    parser = argparse.ArgumentParser(description="Validate extraction accuracy")
    parser.add_argument("command", choices=["grounding", "sample", "score"])
    {"grounding": grounding, "sample": sample, "score": score}[parser.parse_args().command]()


if __name__ == "__main__":
    main()
